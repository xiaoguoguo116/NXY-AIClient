"""Synthetic files for packaged Office Agent verification; never uses user files."""
import hashlib
import json
import sys
from pathlib import Path
from openpyxl import Workbook, load_workbook
from docx import Document

mode, directory = sys.argv[1:]
root = Path(directory).resolve()
root.mkdir(parents=True, exist_ok=True)

def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

if mode == "create":
    book = Workbook()
    book.active.title = "明细"
    for row in [("部门", "金额"), (" 支付 ", 100.25), ("支付", 200.75), ("支付", 200.75), ("风控", 50)]:
        book.active.append(row)
    book.save(root / "input.xlsx")
    book.close()
    doc = Document()
    doc.add_heading("内部制度测试稿", 0)
    paragraph = doc.add_paragraph()
    paragraph.add_run("每").bold = True
    paragraph.add_run("月检查一次，并记录结果。")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "责任部门"
    table.cell(0, 1).text = "原业务部"
    doc.save(root / "input.docx")
    (root / "original-hashes.json").write_text(json.dumps({name: digest(root / name) for name in ["input.xlsx", "input.docx"]}), encoding="utf-8")
    print("Created synthetic Excel and Word inputs")
elif mode == "verify":
    hashes = json.loads((root / "original-hashes.json").read_text(encoding="utf-8"))
    assert all(digest(root / name) == value for name, value in hashes.items()), "Original files changed"
    book = load_workbook(root / "result.xlsx")
    assert list(book["Summary"].values) == [("部门", "sum"), ("支付", 301), ("风控", 50)]
    assert book["明细"].max_row == 4
    assert book["明细"]["A2"].value == "支付"
    book.close()
    doc = Document(root / "result.docx")
    assert doc.paragraphs[1].text == "每周检查一次，并记录结果。"
    assert doc.paragraphs[1].runs[0].bold
    assert doc.tables[0].cell(0, 1).text == "合规部"
    assert doc.paragraphs[-1].text == "本文件为修订草稿，需人工审核。"
    print("PASS: Excel trim/deduplicate/department sums (301 and 50), Word cross-run/table replacements and draft note; original files unchanged.")
else:
    raise ValueError("Unknown fixture operation")
