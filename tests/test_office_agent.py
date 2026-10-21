"""Offline tests using real Office files and a local mock model server.

Run: python -m unittest discover -s tests -p test_office_agent.py -v
Requires the same packages as the built-in office environment.
"""
import importlib.util
import json
import os
import queue
import subprocess
import sys
import tempfile
import threading
import unittest
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from openpyxl import Workbook, load_workbook
from docx import Document
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
AGENT = ROOT / "src-tauri/src/work/office_agent.py"
spec = importlib.util.spec_from_file_location("office_tools", AGENT.with_name("office_tools.py"))
office = importlib.util.module_from_spec(spec)
spec.loader.exec_module(office)


class OfficeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.tools = office.OfficeTools(self.root)
        book = Workbook()
        sheet = book.active
        sheet.title = "销售"
        for row in [["部门", "金额"], [" A ", 10], ["B", 20], [" A ", 10], ["B", 30]]:
            sheet.append(row)
        book.save(self.root / "input.xlsx")
        self.original = (self.root / "input.xlsx").read_bytes()

    def test_cleanup_summary_format_chart_and_formula(self):
        result = self.tools.edit_sheet(path="input.xlsx", output="result.xlsx", operations=[
            {"type": "trim"}, {"type": "deduplicate", "columns": [1, 2]},
            {"type": "sort", "column": 2, "descending": True},
            {"type": "summary", "group_column": 1, "value_column": 2, "target_sheet": "统计"},
            {"type": "format", "range": "A1:B4", "bold": True, "number_format": "0.00"},
            {"type": "chart", "range": "A1:B4", "title": "销售"},
            {"type": "set_cells", "cells": {"D1": {"formula": "=SUM(B2:B4)"}, "E1": "=literal"}}
        ])
        book = load_workbook(result["artifact"]["path"])
        self.addCleanup(book.close)
        self.assertEqual(list(book["统计"].values), [("部门", "sum"), ("B", 50), ("A", 10)])
        self.assertEqual(book["销售"]["B2"].value, 30)
        self.assertTrue(book["销售"]["A1"].font.bold)
        self.assertEqual(book["销售"]["D1"].data_type, "f")
        self.assertEqual(book["销售"]["E1"].data_type, "s")
        self.assertEqual(len(book["销售"]._charts), 1)
        self.assertFalse(result["formulas_evaluated"])
        self.assertEqual((self.root / "input.xlsx").read_bytes(), self.original)

    def test_csv_numeric_summary_and_pagination(self):
        (self.root / "数据.csv").write_text("部门,金额\nA,10\nB,20\nA,30\n", encoding="utf-8-sig")
        preview = self.tools.read_sheet("数据.csv", start_row=2, limit=1)
        self.assertEqual(preview["rows"], [("A", "10")])
        self.assertEqual(preview["next_row"], 3)
        self.tools.edit_sheet(path="数据.csv", output="csv-result.xlsx", operations=[
            {"type": "summary", "group_column": 1, "value_column": 2, "aggregate": "average"}])
        book = load_workbook(self.root / "csv-result.xlsx")
        self.addCleanup(book.close)
        self.assertEqual(list(book["Summary"].values), [("部门", "average"), ("A", 20), ("B", 20)])

    def test_decimal_summary_and_numeric_csv_sort(self):
        (self.root / "input.csv").write_text("group,value\nA,0.1\nA,0.2\nB,10\nB,2\n", encoding="utf-8")
        self.tools.edit_sheet(path="input.csv", output="decimal.xlsx", operations=[
            {"type": "sort", "column": 2, "numeric": True},
            {"type": "summary", "group_column": 1, "value_column": 2}])
        book = load_workbook(self.root / "decimal.xlsx")
        self.addCleanup(book.close)
        self.assertEqual(book.active["B4"].value, "2")
        self.assertEqual(book["Summary"]["B2"].value, 0.3)

    @unittest.skipUnless(os.name == "nt", "Windows canonical paths")
    def test_windows_extended_paths_use_same_scope(self):
        tools = office.OfficeTools("\\\\?\\" + str(self.root.resolve()))
        self.assertEqual(tools.read_sheet(str(self.root / "input.xlsx"))["sheet"], "销售")

    def test_pdf_text_extraction_and_text_pagination(self):
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        writer = PdfWriter()
        page = writer.add_blank_page(width=200, height=200)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 10 100 Td (Office test) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        writer.write(self.root / "input.pdf")
        result = self.tools.read_file("input.pdf", limit=6)
        self.assertEqual(result["text"], "Office")
        self.assertEqual(result["next_offset"], 6)

    def test_docx_cross_run_replacement(self):
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("旧").bold = True
        p.add_run("公司名称")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "旧公司名称"
        doc.save(self.root / "input.docx")
        result = self.tools.write_document(path="input.docx", output="result.docx",
                                          replacements=[{"find": "旧公司", "replace": "新公司"}],
                                          paragraphs=[{"text": "附注", "style": "Heading 1"}])
        doc = Document(result["artifact"]["path"])
        self.assertEqual(doc.paragraphs[0].text, "新公司名称")
        self.assertTrue(doc.paragraphs[0].runs[0].bold)
        self.assertEqual(doc.tables[0].cell(0, 0).text, "新公司名称")
        self.assertEqual(result["replacements"], 2)

    def test_image_conversion(self):
        Image.new("RGBA", (100, 80), (255, 0, 0, 100)).save(self.root / "input.png")
        self.tools.process_image(path="input.png", output="result.jpg", width=50, crop=[0, 0, 80, 80])
        with Image.open(self.root / "result.jpg") as image:
            self.assertEqual(image.size, (50, 50))
            self.assertEqual(image.mode, "RGB")
        with self.assertRaises(ValueError):
            self.tools.process_image(path="input.png", output="bad.jpg", width=0)

    def test_scope_and_no_overwrite(self):
        with tempfile.TemporaryDirectory() as external:
            other = Path(external) / "other.txt"
            other.write_text("read only", encoding="utf-8")
            with self.assertRaises(ValueError):
                self.tools.read_file(str(other))
            tools = office.OfficeTools(self.root, [other])
            self.assertEqual(tools.read_file(str(other))["text"], "read only")
            with self.assertRaises(ValueError):
                tools.write_text(str(Path(external) / "result.txt"), "no")
        with self.assertRaises(ValueError):
            self.tools.edit_sheet(path="input.xlsx", output="input.xlsx", operations=[])
        with self.assertRaises(ValueError):
            self.tools.write_text("../escaped.txt", "no")
        self.assertEqual((self.root / "input.xlsx").read_bytes(), self.original)

    def test_symlink_escape(self):
        with tempfile.TemporaryDirectory() as external:
            link = self.root / "link"
            try:
                link.symlink_to(external, target_is_directory=True)
            except OSError:
                self.skipTest("Symlink privilege unavailable")
            with self.assertRaises(ValueError):
                self.tools.write_text("link/escaped.txt", "no")

    def test_formula_reorder_refused(self):
        self.tools.edit_sheet(path="input.xlsx", output="formula.xlsx",
                              operations=[{"type": "set_cells", "cells": {"C2": {"formula": "=B2"}}}])
        for op in [{"type": "sort", "column": 2}, {"type": "deduplicate"}]:
            with self.assertRaises(ValueError):
                self.tools.edit_sheet(path="formula.xlsx", output="bad.xlsx", operations=[op])
        self.assertFalse((self.root / "bad.xlsx").exists())

    def test_cancelled_publish_and_failed_tool_leave_no_files(self):
        tools = office.OfficeTools(self.root, cancelled=lambda: True)
        with self.assertRaises(InterruptedError):
            tools.write_text("cancelled.txt", "no")
        with self.assertRaises(ValueError):
            self.tools.edit_sheet(output="bad.xlsx", operations=[{"type": "unknown"}])
        self.assertFalse((self.root / "cancelled.txt").exists())
        self.assertFalse((self.root / "bad.xlsx").exists())
        self.assertFalse(list(self.root.glob(".em-work-*")))

    def test_unknown_tool_and_legacy_format(self):
        with self.assertRaises(ValueError):
            self.tools.execute("execute_shell", {"command": "anything"})
        (self.root / "old.xls").write_bytes(b"old")
        with self.assertRaises(ValueError):
            self.tools.read_sheet("old.xls")

    def test_text_copy_and_csv_injection(self):
        self.tools.write_text("note.txt", "你好")
        self.tools.copy_file("note.txt", "copy.txt")
        self.assertEqual(self.tools.read_file("copy.txt")["text"], "你好")
        with self.assertRaises(ValueError):
            self.tools.copy_file("note.txt", "evil.cmd")
        self.tools.edit_sheet(output="safe.csv", operations=[
            {"type": "append_rows", "rows": [["name"], ["=HYPERLINK(\"bad\")"]]}])
        self.assertIn("'=HYPERLINK", (self.root / "safe.csv").read_text(encoding="utf-8-sig"))


class ProtocolTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.requests = []
        self.responses = queue.Queue()
        case = self

        class Handler(BaseHTTPRequestHandler):
            def do_POST(self):
                case.requests.append(json.loads(self.rfile.read(int(self.headers["Content-Length"]))))
                response = case.responses.get(timeout=10)
                if callable(response):
                    response = response()
                if isinstance(response, int):
                    self.send_response(response)
                    self.end_headers()
                else:
                    body = json.dumps(response).encode()
                    self.send_response(200)
                    self.send_header("Content-Type", "application/json")
                    self.send_header("Content-Length", str(len(body)))
                    self.end_headers()
                    self.wfile.write(body)

            def log_message(self, *args):
                pass

        self.server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        self.server.daemon_threads = True
        threading.Thread(target=self.server.serve_forever, daemon=True).start()
        self.addCleanup(self.server.server_close)
        self.addCleanup(self.server.shutdown)

    def start(self, kind="openai"):
        env = {**os.environ, "TEST_KEY": "local-test",
               "EM_OFFICE_CONFIG": json.dumps({"baseUrl": f"http://127.0.0.1:{self.server.server_port}/v1",
                                              "kind": kind, "apiKeyEnv": "TEST_KEY", "model": "mock"})}
        self.process = subprocess.Popen([sys.executable, "-I", "-u", str(AGENT)], env=env,
                                        stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                                        text=True, encoding="utf-8")
        self.addCleanup(self.stop)
        self.messages = queue.Queue()
        self.seen = []
        self.stderr = []
        def read_errors():
            for line in self.process.stderr:
                self.stderr.append(line)
        threading.Thread(target=read_errors, daemon=True).start()
        def reader():
            for line in self.process.stdout:
                self.messages.put(json.loads(line))
        threading.Thread(target=reader, daemon=True).start()
        self.send(1, "initialize", {})
        self.assertEqual(self.until(lambda m: m.get("id") == 1)["result"]["protocolVersion"], 1)
        self.send(2, "session/new", {"cwd": str(self.root)})
        self.sid = self.until(lambda m: m.get("id") == 2)["result"]["sessionId"]

    def stop(self):
        self.process.terminate()
        self.process.wait(timeout=10)
        for stream in [self.process.stdin, self.process.stdout, self.process.stderr]:
            stream.close()

    def send(self, request_id, method, params):
        self.process.stdin.write(json.dumps({"jsonrpc": "2.0", "id": request_id, "method": method, "params": params}) + "\n")
        self.process.stdin.flush()

    def until(self, predicate):
        while True:
            try:
                message = self.messages.get(timeout=15)
            except queue.Empty:
                self.fail(f"Timed out waiting for protocol message. Recent messages: {self.seen[-2:]}; stderr: {self.stderr}")
            self.seen.append(message)
            if predicate(message):
                return message

    def prompt(self):
        self.send(3, "session/prompt", {"sessionId": self.sid, "prompt": [{"type": "text", "text": "生成表格"}]})

    def tool_response(self, name="edit_sheet", arguments=None):
        return {"choices": [{"message": {"role": "assistant", "content": None, "tool_calls": [{
            "id": "tool-1", "type": "function", "function": {"name": name, "arguments": json.dumps(arguments or {
                "output": "result.xlsx", "operations": [{"type": "append_rows", "rows": [["金额"], [123]]}]})}}]}}]}

    def final_response(self):
        return {"choices": [{"message": {"role": "assistant", "content": "处理结束"}}]}

    def approve(self, option):
        message = self.until(lambda m: m.get("method") == "session/request_permission")
        self.process.stdin.write(json.dumps({"id": message["id"], "result": {"outcome": {
            "outcome": "selected", "optionId": option}}}) + "\n")
        self.process.stdin.flush()

    def test_openai_tool_loop_approval_artifact(self):
        self.start()
        self.responses.put(self.tool_response())
        self.responses.put(self.final_response())
        self.prompt()
        self.approve("allow_once")
        artifact = self.until(lambda m: m.get("params", {}).get("update", {}).get("sessionUpdate") == "artifact")
        self.assertTrue(Path(artifact["params"]["update"]["path"]).is_file())
        self.assertEqual(self.until(lambda m: m.get("id") == 3)["result"]["stopReason"], "end_turn")
        self.assertEqual(self.requests[1]["messages"][-1]["role"], "tool")
        self.assertIn("artifact", self.requests[1]["messages"][-1]["content"])

    def test_permission_denial_writes_nothing(self):
        self.start()
        self.responses.put(self.tool_response())
        self.responses.put(self.final_response())
        self.prompt()
        self.approve("reject_once")
        self.until(lambda m: m.get("id") == 3)
        self.assertFalse((self.root / "result.xlsx").exists())
        self.assertIn("denied", self.requests[1]["messages"][-1]["content"])

    def test_denial_cannot_be_bypassed_with_another_tool(self):
        self.start()
        self.responses.put(self.tool_response())
        self.responses.put(self.tool_response("write_text", {"output": "bypass.txt", "text": "no"}))
        self.responses.put(self.final_response())
        self.prompt()
        self.approve("reject_once")
        self.until(lambda m: m.get("id") == 3)
        self.assertFalse(list(self.root.iterdir()))
        self.assertIn("denied writes", self.requests[2]["messages"][-1]["content"])

    def test_cancel_while_waiting_for_permission(self):
        self.start()
        self.responses.put(self.tool_response())
        self.prompt()
        self.until(lambda m: m.get("method") == "session/request_permission")
        self.send(4, "session/cancel", {"sessionId": self.sid})
        result = self.until(lambda m: m.get("id") == 3)
        self.assertEqual(result["result"]["stopReason"], "cancelled")
        self.assertFalse((self.root / "result.xlsx").exists())

    def test_cancel_during_model_request_never_executes_tools(self):
        self.start()
        entered, release = threading.Event(), threading.Event()
        self.addCleanup(release.set)
        def delayed():
            entered.set()
            release.wait(timeout=8)
            return self.tool_response()
        self.responses.put(delayed)
        self.prompt()
        self.assertTrue(entered.wait(timeout=5))
        self.send(4, "session/cancel", {"sessionId": self.sid})
        self.until(lambda m: m.get("id") == 4)
        release.set()
        self.assertEqual(self.until(lambda m: m.get("id") == 3)["result"]["stopReason"], "cancelled")
        self.assertFalse(list(self.root.iterdir()))

    def test_anthropic_tool_loop(self):
        self.start("anthropic")
        self.responses.put({"content": [{"type": "tool_use", "id": "anthropic-1", "name": "write_text",
                                        "input": {"output": "note.txt", "text": "真实结果"}}]})
        self.responses.put({"content": [{"type": "text", "text": "完成"}]})
        self.prompt()
        self.approve("allow_once")
        self.until(lambda m: m.get("id") == 3)
        self.assertEqual((self.root / "note.txt").read_text(encoding="utf-8"), "真实结果")
        self.assertEqual(self.requests[1]["messages"][-1]["content"][0]["type"], "tool_result")

    def test_http_failure_is_error(self):
        self.start()
        self.responses.put(401)
        self.prompt()
        self.assertIn("HTTP 401", self.until(lambda m: m.get("id") == 3)["error"]["message"])
        self.assertFalse(list(self.root.iterdir()))

    def test_model_cannot_authorize_arbitrary_attachment(self):
        self.start()
        with tempfile.TemporaryDirectory() as external:
            target = Path(external) / "secret.txt"
            target.write_text("outside", encoding="utf-8")
            self.responses.put(self.tool_response("read_file", {"path": str(target)}))
            self.responses.put(self.final_response())
            self.prompt()
            self.until(lambda m: m.get("id") == 3)
            self.assertIn("outside the workspace", self.requests[1]["messages"][-1]["content"])

    def test_explicit_attachment_grants_read_access(self):
        self.start()
        with tempfile.TemporaryDirectory() as external:
            target = Path(external) / "attached.txt"
            target.write_text("attached data", encoding="utf-8")
            self.responses.put(self.tool_response("read_file", {"path": str(target)}))
            self.responses.put(self.final_response())
            self.send(3, "session/prompt", {"sessionId": self.sid,
                "prompt": [{"type": "text", "text": "read attachment"}], "_meta": {"emAttachments": [str(target)]}})
            self.until(lambda m: m.get("id") == 3)
            self.assertIn("attached data", self.requests[1]["messages"][-1]["content"])


if __name__ == "__main__":
    unittest.main()
